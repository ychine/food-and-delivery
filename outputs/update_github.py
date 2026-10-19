from docx import Document
from docx.oxml import OxmlElement
from docx.text.paragraph import Paragraph
from docx.shared import Pt
path='outputs/Food_Delivery_Group_Procedures.docx'
doc=Document(path)
for p in doc.paragraphs:
 if p.text.startswith('This guide assigns'):
  p.text='This guide assigns the project procedures to Groups 1 through 7. Each group uses its own computer and retrieves the shared project files from GitHub. Use Docker to build and run the services, AWS to host them, and Postman to test the APIs. Commands below use PowerShell. Start Docker before building or running containers.'
 if p.text.startswith('Prepare the Restaurant Service'):
  p.text='Prepare the Restaurant Service on port 3001 and the Menu Service on port 3002, with a Dockerfile for each service. Commit and push the service files, Dockerfiles, and docker-compose.yml to the shared GitHub repository. Other groups clone the repository once and pull the latest changes before continuing. Run the following commands from the main Food_Delivery folder.'
 if p.text.startswith('Use the computer where Group 1'):
  p.text='On your own computer, open the local clone of the shared GitHub repository and pull the latest changes. GitHub shares the source files and Dockerfiles; build the Docker images locally before tagging and pushing them. Configure AWS CLI and authenticate Docker with ECR on this computer using the Group 3 procedure.'
  anchor=p
  for text,iscode in [('git pull\ndocker build -t restaurant-service ./restaurant-service\ndocker build -t menu-service ./menu-service',True),('In ECR, open restaurant-service, select View push commands, and run the displayed tag and push commands for restaurant-service:latest. Repeat for menu-service:latest in the menu-service repository.',False)]:
   e=OxmlElement('w:p'); anchor._p.addnext(e); q=Paragraph(e,anchor._parent); r=q.add_run(text)
   if iscode: r.font.name='Consolas'; r.font.size=Pt(10)
   anchor=q
 if p.text.startswith('Configure the AWS CLI on the computer'):
  p.text='Configure AWS CLI on each computer that will perform AWS CLI operations or upload images. GitHub provides the shared project files; each group uses its own authorized AWS credentials and selects the region containing the project repositories.'
doc.save(path)
print('Updated GitHub workflow in document')
