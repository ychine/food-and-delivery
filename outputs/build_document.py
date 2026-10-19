from docx import Document
from docx.shared import Inches, Pt, RGBColor
from pathlib import Path
out=Path('outputs')
doc=Document()
s=doc.sections[0]; s.top_margin=s.bottom_margin=Inches(.65); s.left_margin=s.right_margin=Inches(.75)
normal=doc.styles['Normal']; normal.font.name='Calibri'; normal.font.size=Pt(11); normal.paragraph_format.space_after=Pt(6)
for n in ['Title','Heading 1','Heading 2']:
 doc.styles[n].font.color.rgb=RGBColor(0,0,0)
doc.styles['Title'].font.size=Pt(24)
doc.styles['Heading 1'].font.size=Pt(15)
doc.styles['Heading 1'].paragraph_format.space_before=Pt(12)
def p(t): doc.add_paragraph(t)
def h(t,steps):
 doc.add_heading(t,1); q=doc.add_paragraph(steps); q.runs[0].italic=True

def code(t):
 q=doc.add_paragraph(); q.paragraph_format.space_before=Pt(3); q.paragraph_format.space_after=Pt(8)
 for i,line in enumerate(t.split('\n')):
  r=q.add_run(('\n' if i else '')+line); r.font.name='Consolas'; r.font.size=Pt(10)
def bullet(t): doc.add_paragraph(t,style='List Bullet')
doc.add_heading('Food Delivery API Group Procedures',0)
p('This guide assigns the project procedures to Groups 1 through 7. Use Docker to build and run the services, AWS to host them, and Postman to test the APIs. Commands below use PowerShell. Start Docker before building or running containers.')
h('Group 1 Services and Docker Images','Steps 1 to 5')
p('Prepare the Restaurant Service on port 3001 and the Menu Service on port 3002, with a Dockerfile for each service. Run the following commands from the main Food_Delivery folder.')
code('node --version\nnpm --version\ndocker --version\ndocker build -t restaurant-service ./restaurant-service\ndocker build -t menu-service ./menu-service')
p('Run both services locally with Docker Compose:')
code('docker compose up --build -d')
p('In Postman, select GET, enter each URL below, and click Send. The third request tests communication between the two services.')
code('http://localhost:3001/restaurants\nhttp://localhost:3002/menu\nhttp://localhost:3001/restaurants/1/menu')
h('Group 2 ECR Repositories','Step 6')
p('In the AWS Console, open Elastic Container Registry, select Private repositories, and choose Create repository. Create restaurant-service and menu-service. These repositories store the Docker images uploaded by Group 4. If they already exist, use the existing repositories.')
doc.add_page_break()
h('Group 3 AWS CLI and Docker Authentication','Steps 7 and 8')
p('Configure the AWS CLI on the computer that will upload the images. Use authorized AWS credentials and select the region containing the project repositories.')
code('aws --version\naws configure\naws sts get-caller-identity')
p('In AWS ECR, open restaurant-service and select View push commands. Copy and run the login command provided there. It already contains the correct AWS account ID and region. Use the Windows command instructions when available. Keep AWS secret keys private.')
h('Group 4 Tag and Push Images','Steps 9 and 10')
p('Use the computer where Group 1 built the Docker images. In ECR, open restaurant-service, select View push commands, and run the displayed tag and push commands for restaurant-service:latest. Repeat for menu-service:latest in the menu-service repository.')
p('After uploading, confirm that both ECR repositories contain an image with the latest tag. The account-specific tag and push commands should be copied directly from ECR.')
h('Group 5 ECS Cluster','Step 11')
p('In the AWS Console, open Elastic Container Service, select Clusters, and choose Create cluster. Create the cluster for the Restaurant and Menu services using AWS Fargate. This procedure is completed in the AWS Console.')
h('Group 6 Task Definitions and Fargate','Steps 12 to 14')
p('In ECS, open Task definitions and choose Create new task definition. Create separate task definitions for the two services:')
bullet('Restaurant Service: use the Restaurant ECR image and container port 3001.')
bullet('Menu Service: use the Menu ECR image and container port 3002.')
p('Configure Fargate compatibility, CPU, memory, and a task execution role with access to ECR. In the Restaurant task, set MENU_SERVICE_URL to the reachable Menu Service address. If AWS Cloud Map is configured, use the actual Menu DNS name. A public Menu IP must be updated if the task IP changes.')
doc.add_page_break()
h('Group 7 Networking and Deployment','Steps 15 and 16')
p('Within the ECS cluster, create a service for each task definition. Configure the VPC, subnets, Fargate networking, and security groups. For direct access through the public links, enable a public IP and use subnets with internet connectivity.')
p('Configure the following inbound security group rules:')
bullet('Restaurant Service: Custom TCP, port 3001, source Anywhere-IPv4 (0.0.0.0/0).')
bullet('Menu Service: Custom TCP, port 3002, source Anywhere-IPv4 (0.0.0.0/0).')
p('Keep the existing Menu rule whose source is a security group ID beginning with sg-. Click Save rules for each security group. Replacing the individual IP address ending in /32 with Anywhere-IPv4 allows classmates and other internet users to reach the specified service ports.')
h('Group 1 Final Testing with Postman','Step 18')
p('Create a Postman collection named Food Delivery API. Add the following requests, select GET for each, and click Send.')
for title,url in [
 ('Restaurant status','http://3.26.198.35:3001/'),
 ('Restaurant list','http://3.26.198.35:3001/restaurants'),
 ('Menu status','http://16.176.192.32:3002/'),
 ('Menu list','http://16.176.192.32:3002/menu'),
 ('Restaurant and Menu integration','http://3.26.198.35:3001/restaurants/1/menu')]:
 q=doc.add_paragraph(); q.add_run(title).bold=True
 code(url)
p('Expected result: 200 OK with a JSON response. A 503 response from the integration request means the Restaurant Service could not complete its request to the Menu Service; check the Menu address and connectivity between the services.')
p('Save the requests and screenshots of successful results for documentation. If ECS replaces a task and its public IP changes, update the corresponding Postman URLs.')
doc.core_properties.title='Food Delivery API Group Procedures'
doc.core_properties.subject='Docker AWS and Postman procedures by group'
doc.save(out/'Food_Delivery_Group_Procedures.docx')
print((out/'Food_Delivery_Group_Procedures.docx').resolve())
